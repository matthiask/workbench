import datetime as dt
from decimal import Decimal
from io import BytesIO

from django.core.exceptions import ValidationError
from django.test import RequestFactory, TestCase
from django.urls import reverse
from django.utils.translation import deactivate_all
from openpyxl import load_workbook

from workbench import factories
from workbench.accounts.models import User
from workbench.planning import reporting
from workbench.planning.forms import MilestoneForm, PlannedWorkSearchForm
from workbench.planning.models import ExternalWork, PlannedWork
from workbench.tools.validation import in_days, monday


date_range = [in_days(-14), in_days(400)]


class PlanningTest(TestCase):
    def setUp(self):
        deactivate_all()

    def test_no_monday(self):
        """Non-mondays are rejected from PlannedWork.weeks"""
        pw = factories.PlannedWorkFactory.create(weeks=[dt.date(2020, 6, 21)])

        msg = ["Only mondays allowed, but field contains 21.06.2020."]

        with self.assertRaises(ValidationError) as cm:
            pw.clean_fields(exclude=["weeks"])
        self.assertEqual(list(cm.exception), msg)

        pw.weeks = [dt.date(2020, 6, 22)]
        pw.full_clean()  # Does not raise
        self.assertEqual(pw.pretty_from_until, "22.06.2020 – 28.06.2020")

    def test_reporting_smoke(self):
        """Smoke test the planned work report"""
        pw = factories.PlannedWorkFactory.create(weeks=[monday()])

        factories.EmploymentFactory.create(user=pw.user, date_from=dt.date(2020, 1, 1))
        service = factories.ServiceFactory.create(project=pw.project)
        factories.LoggedHoursFactory.create(rendered_by=pw.user, service=service)
        factories.AbsenceFactory.create(user=pw.user)

        report = reporting.user_planning(pw.user, date_range)
        self.assertAlmostEqual(sum(report["by_week"]), Decimal(28))
        self.assertEqual(len(report["projects_offers"]), 1)

        report = reporting.project_planning(pw.project)
        self.assertAlmostEqual(sum(report["by_week"]), Decimal(28))
        self.assertEqual(len(report["projects_offers"]), 1)

        team = factories.TeamFactory.create()
        team.members.add(pw.user)
        report = reporting.team_planning(team, date_range)
        self.assertAlmostEqual(sum(report["by_week"]), Decimal(28))
        self.assertEqual(len(report["projects_offers"]), 1)

        pw2 = factories.PlannedWorkFactory.create(
            project=pw.project,
            user=pw.user,
            weeks=[monday()],
        )
        report = reporting.user_planning(pw.user, date_range)
        self.assertAlmostEqual(sum(report["by_week"]), Decimal(48))
        self.assertEqual(len(report["projects_offers"]), 1)

        work_list = report["projects_offers"][0]["offers"][0]["work_list"]
        self.assertEqual(len(work_list), 2)
        self.assertEqual(work_list[0]["work"]["id"], pw2.id)
        self.assertEqual(work_list[1]["work"]["id"], pw.id)

        report = reporting.planning_vs_logbook(date_range, users=User.objects.all())
        self.assertAlmostEqual(report["logged"], Decimal("1.0"))
        self.assertAlmostEqual(report["planned"], Decimal("40.0"))
        # Exactly one customer
        (c,) = report["per_customer"]
        self.assertEqual(len(c["per_week"]), 1)

    def test_milestone_color(self):
        """The milestone color overrides the service type color of its work"""
        milestone = factories.MilestoneFactory.create(
            date=monday() + dt.timedelta(days=3), color="#123456"
        )
        pw = factories.PlannedWorkFactory.create(
            project=milestone.project, milestone=milestone, weeks=[monday()]
        )
        other = factories.PlannedWorkFactory.create(
            project=milestone.project, user=pw.user, weeks=[monday()]
        )

        report = reporting.project_planning(milestone.project)
        (project,) = report["projects_offers"]
        (ms,) = project["project"]["milestones"]
        (offer,) = project["offers"]
        colors = {
            entry["work"]["id"]: entry["work"]["color"]
            for entry in ms["work_list"] + offer["work_list"]
        }
        self.assertEqual(colors[pw.id], "#123456")
        self.assertEqual(colors[other.id], "")
        self.assertEqual(project["project"]["milestones"][0]["color"], "#123456")

    def test_work_below_milestones(self):
        """Work belonging to a milestone is shown below it, the rest is
        grouped by offer"""
        milestone = factories.MilestoneFactory.create(date=in_days(10))
        project = milestone.project
        offer = factories.OfferFactory.create(project=project)
        pw = factories.PlannedWorkFactory.create(
            project=project,
            offer=offer,
            milestone=milestone,
            weeks=[monday()],
            planned_hours=10,
        )
        other = factories.PlannedWorkFactory.create(
            project=project,
            offer=offer,
            user=pw.user,
            weeks=[monday() + dt.timedelta(days=7)],
            planned_hours=5,
        )
        ew = ExternalWork.objects.create(
            project=project,
            milestone=milestone,
            provided_by=factories.OrganizationFactory.create(),
            created_by=pw.user,
            title="Printing",
            weeks=[monday()],
        )

        (record,) = reporting.project_planning(project)["projects_offers"]
        (ms,) = record["project"]["milestones"]
        self.assertEqual([wl["work"]["id"] for wl in ms["work_list"]], [pw.id])
        self.assertEqual([e["id"] for e in ms["external_work"]], [ew.id])
        self.assertEqual(ms["work_list"][0]["work"]["offer"]["title"], offer.title)
        self.assertIn("absences", ms["work_list"][0])

        (offer_record,) = record["offers"]
        self.assertEqual(
            [wl["work"]["id"] for wl in offer_record["work_list"]], [other.id]
        )
        self.assertEqual(offer_record["offer"]["planned_hours"], 5)
        self.assertEqual(record["project"]["planned_hours"], 15)
        self.assertIsNone(record["external_work"])

        # The Gantt chart lists the work below its milestone as well
        self.client.force_login(pw.user)
        response = self.client.get(
            reverse("projects_project_planning_gantt_xlsx", kwargs={"pk": project.pk})
        )
        ws = load_workbook(BytesIO(response.content))["Gantt Chart"]
        self.assertEqual(
            [row[1] for row in ws.iter_rows(min_row=2, max_col=2, values_only=True)],
            [milestone.title, pw.title, ew.title, other.title],
        )

        # Offers without work outside milestones aren't shown
        other.delete()
        (record,) = reporting.project_planning(project)["projects_offers"]
        self.assertIsNone(record["offers"])
        self.assertEqual(record["project"]["planned_hours"], 10)

    def test_milestone_color_choices(self):
        """Milestone colors are chosen from swatches, not typed as hex codes"""
        milestone = factories.MilestoneFactory.create(date=in_days(10), color="#123456")
        request = RequestFactory().get("/")
        request.user = milestone.project.owned_by

        def form(color, instance=milestone):
            return MilestoneForm(
                {
                    "modal-date": milestone.date.isoformat(),
                    "modal-title": "Launch",
                    "modal-estimated_total_hours": "0",
                    "modal-color": color,
                },
                instance=instance,
                request=request,
            )

        # Colors entered before are kept as an additional choice
        self.assertIn('value="#123456"', str(form("#123456")["color"]))
        self.assertTrue(form("#3b5bdb").is_valid())
        self.assertTrue(form("").is_valid())
        self.assertFalse(form("rot").is_valid())
        other = factories.MilestoneFactory.create(date=in_days(10))
        self.assertFalse(form("#123456", instance=other).is_valid())

    def test_planning_search_forms(self):
        """Planning request search form branch test"""

        project = factories.ProjectFactory.create()
        pw = factories.PlannedWorkFactory.create(weeks=[dt.date(2020, 6, 22)])

        rf = RequestFactory()

        for form_class, model, instances in [
            (PlannedWorkSearchForm, PlannedWork, [pw]),
        ]:
            req = rf.get("/")
            req.user = project.owned_by
            form = form_class(req.GET, request=req)
            self.assertTrue(form.is_valid())
            queryset = form.filter(model.objects.all())
            self.assertEqual(list(queryset), instances)

            req = rf.get(f"/?project={project.pk}")
            req.user = project.owned_by
            form = form_class(req.GET, request=req)
            self.assertTrue(form.is_valid())
            queryset = form.filter(model.objects.all())
            self.assertEqual(list(queryset), [])

    def test_planned_work_with_offer(self):
        """Planned work with offer preselection"""
        offer = factories.OfferFactory.create()
        self.client.force_login(offer.owned_by)

        response = self.client.get(
            offer.project.urls["creatework"] + f"?offer={offer.pk}"
        )
        self.assertContains(
            response,
            f'<option value="{offer.pk}" selected>{offer}</option>',
            html=True,
        )

    def test_planned_work_crud(self):
        """Create, update and delete planned work"""
        service_types = factories.service_types()

        project = factories.ProjectFactory.create()
        self.client.force_login(project.owned_by)

        response = self.client.get(project.urls["creatework"] + "?request=bla")
        self.assertEqual(response.status_code, 200)  # No crash

        response = self.client.post(
            project.urls["creatework"],
            {
                "modal-user": project.owned_by.id,
                "modal-title": "bla",
                "modal-planned_hours": 50,
                "modal-weeks": [monday().isoformat()],
                "modal-service_type": service_types.consulting.pk,
            },
            headers={"x-requested-with": "XMLHttpRequest"},
        )
        self.assertEqual(response.status_code, 201)

        pw = PlannedWork.objects.get()
        response = self.client.post(
            pw.urls["update"],
            {
                "modal-user": project.owned_by.id,
                "modal-title": "bla",
                "modal-planned_hours": 50,
                "modal-weeks": [monday().isoformat()],
                "modal-service_type": service_types.consulting.pk,
            },
            headers={"x-requested-with": "XMLHttpRequest"},
        )
        self.assertEqual(response.status_code, 202)

        response = self.client.post(
            project.urls["creatework"],
            {
                "modal-user": project.owned_by.id,
                "modal-title": "bla",
                "modal-planned_hours": 50,
                "modal-weeks": [
                    monday().isoformat(),
                    (monday() + dt.timedelta(days=7)).isoformat(),
                ],
                "modal-service_type": service_types.consulting.pk,
            },
            headers={"x-requested-with": "XMLHttpRequest"},
        )
        # print(response, response.content.decode("utf-8"))
        self.assertEqual(response.status_code, 201)

        # response = self.client.post(
        #     project.urls["creatework"],
        #     {
        #         "modal-user": project.owned_by.id,
        #         "modal-title": "bla",
        #         "modal-planned_hours": 50,
        #         "modal-weeks": [
        #             monday().isoformat(),
        #             (monday() + dt.timedelta(days=14)).isoformat(),
        #         ],
        #     },
        #     HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        # )
        # self.assertContains(response, "weeks-outside-request")

    def test_planning_views(self):
        """The planning views do not crash"""
        user = factories.UserFactory.create()
        team = factories.TeamFactory.create()
        project = factories.ProjectFactory.create()

        self.client.force_login(user)

        response = self.client.get(user.urls["planning"])
        self.assertContains(response, 'id="planning-data"')

        response = self.client.get(team.urls["planning"])
        self.assertContains(response, 'id="planning-data"')

        response = self.client.get(project.urls["planning"])
        self.assertContains(response, 'id="planning-data"')

    def test_this_week_index(self):
        """this_week_index is None if current week isn't part of the report"""
        pw = factories.PlannedWorkFactory.create(weeks=[dt.date(2020, 6, 22)])

        report = reporting.project_planning(pw.project)
        self.assertIsNone(report["this_week_index"])

    def test_planning_request_notification(self):
        """Receivers of a planning request are notified"""
        # pr = factories.PlanningRequestFactory.create()
        # pr.receivers.add(
        #     factories.UserFactory.create(),
        #     factories.UserFactory.create(),
        #     factories.UserFactory.create(),
        # )
        # self.assertEqual(len(mail.outbox), 1)
        # self.assertIn("New planning request", mail.outbox[0].subject)
        # self.assertEqual(len(mail.outbox[0].to), 3)
        # self.assertEqual(len(mail.outbox[0].reply_to), 4)

    def test_declined_offer_warning(self):
        """Warn when offer is declined"""
        offer = factories.OfferFactory.create(status=factories.Offer.DECLINED)
        self.client.force_login(offer.owned_by)

        response = self.client.post(
            offer.project.urls["creatework"],
            {
                "modal-user": offer.owned_by_id,
                "modal-title": "bla",
                "modal-planned_hours": 50,
                "modal-weeks": [monday().isoformat()],
                "modal-offer": offer.id,
            },
            headers={"x-requested-with": "XMLHttpRequest"},
        )
        self.assertContains(response, 'value="offer-is-declined"')

    def test_planned_work_ranges(self):
        """Range detection works"""
        pw = factories.PlannedWorkFactory.create(
            weeks=[
                dt.date(2020, 6, 22),
                dt.date(2020, 6, 29),
                # dt.date(2020, 7, 6),
                dt.date(2020, 7, 13),
                dt.date(2020, 7, 20),
            ]
        )
        self.assertEqual(
            list(pw.ranges),
            [
                {
                    "from": dt.date(2020, 6, 22),
                    "until": dt.date(2020, 6, 29),
                    "pretty": "22.06.2020 – 05.07.2020",
                },
                {
                    "from": dt.date(2020, 7, 13),
                    "until": dt.date(2020, 7, 20),
                    "pretty": "13.07.2020 – 26.07.2020",
                },
            ],
        )
        self.assertEqual(pw.pretty_planned_hours, "20.0h in 4 weeks (5.0h per week)")

    def test_initialize_form_using_offer(self):
        """Initializing planning requests and planned work using offers"""

        offer = factories.OfferFactory.create(
            title="Testing title",
            description="Testing description",
        )
        self.client.force_login(offer.owned_by)

        response = self.client.get(
            offer.project.urls["creatework"] + f"?plan_offer={offer.pk}"
        )
        self.assertContains(response, 'value="Testing title"')
        self.assertContains(response, "Testing description")

        response = self.client.get(offer.project.urls["creatework"] + "?plan_offer=bla")
        self.assertEqual(response.status_code, 200)  # No crash

    def test_initialize_form_using_service(self):
        """Initializing planning requests and planned work using services"""

        service = factories.ServiceFactory.create(
            title="Testing title",
            description="Testing description",
            effort_hours=20,
        )
        self.client.force_login(service.project.owned_by)

        response = self.client.get(
            service.project.urls["creatework"] + f"?service={service.pk}"
        )
        self.assertContains(response, f'value="{service.project.title}: Testing title"')
        self.assertContains(response, "Testing description")
        self.assertContains(response, 'value="20.0"')

        response = self.client.get(service.project.urls["creatework"] + "?service=bla")
        self.assertEqual(response.status_code, 200)  # No crash

    def test_initialize_form_using_weeks(self):
        """Weeks dragged in the planning and the milestone of the row are
        filled in"""
        milestone = factories.MilestoneFactory.create(date=in_days(60))
        self.client.force_login(milestone.project.owned_by)

        weeks = [monday() + dt.timedelta(days=days) for days in (7, 14)]
        response = self.client.get(
            milestone.project.urls["creatework"]
            + f"?milestone={milestone.pk}&weeks={weeks[0]}&weeks={weeks[1]}"
        )
        form = response.context["form"]
        self.assertEqual(form["weeks"].value(), weeks)
        self.assertEqual(form["milestone"].value(), str(milestone.pk))

        response = self.client.get(
            milestone.project.urls["creatework"] + "?weeks=bla&weeks=2026-13-40"
        )
        self.assertEqual(response.context["form"]["weeks"].value(), [monday()])

    def test_decision_of_open_deal(self):
        """Offers which aren't accepted show when their open deal should be
        decided"""
        offer = factories.OfferFactory.create()
        # The external view only shows work belonging to milestones
        milestone = factories.MilestoneFactory.create(
            project=offer.project, date=in_days(30)
        )
        pw = factories.PlannedWorkFactory.create(
            project=offer.project, offer=offer, milestone=milestone, weeks=[monday()]
        )
        factories.PlannedWorkFactory.create(
            project=offer.project, offer=offer, weeks=[monday()]
        )
        deal = factories.DealFactory.create(decision_expected_on=in_days(-3))
        deal.related_offers.add(offer)

        def decision(**kwargs):
            report = reporting.project_planning(pw.project, **kwargs)
            (project,) = report["projects_offers"]
            if not project["offers"]:
                return None
            (row,) = project["offers"]
            return row["offer"].get("decision_overdue")

        self.assertTrue(decision())
        self.assertIsNone(decision(external_view=True))
        deal.status = deal.ACCEPTED
        deal.closed_on = dt.date.today()
        deal.save()
        self.assertIsNone(decision())
