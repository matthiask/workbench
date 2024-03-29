import sys
from collections import defaultdict

from openpyxl import load_workbook
from xlsxdocument import XLSXDocument


workbooks = [load_workbook(name) for name in sys.argv[1:]]
sheets = [list(wb.worksheets[0].rows) for wb in workbooks]
users = defaultdict(
    lambda: {
        "percentage": 0,
        "gross_margin": [None for _ in sheets],
        "expected_gross_margin": 0,
    }
)

for idx, sheet in enumerate(sheets):
    for row in sheet[4:]:
        users[row[0].value]["percentage"] = row[2].value
        users[row[0].value]["expected_gross_margin"] = row[22].value
        users[row[0].value]["gross_margin"][idx] = row[3].value

xlsx = XLSXDocument()
xlsx.add_sheet("Nutzer_innen")
xlsx.table(
    [sheets[0][1][0].value]
    + [sheet[0][0].value.removeprefix("Squeeze ") for sheet in sheets]
    + [sheets[0][1][2].value, sheets[0][1][22].value],
    [
        [user]
        + data["gross_margin"]
        + [data["percentage"], data["expected_gross_margin"]]
        for user, data in sorted(users.items())
    ],
)
xlsx.workbook.save("merge_squeeze.xlsx")
